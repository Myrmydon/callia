# Nouvelle version du main.py avec interface TTK Bootstrap modernisée, EPUB, compteur et couleur verte personnalisée
import tkinter as tk
from tkinter import filedialog, messagebox
from ttkbootstrap import Window, Style, ttk
from ttkbootstrap.icons import Icon
from reportlab.platypus import SimpleDocTemplate, Paragraph, PageBreak, Spacer
from reportlab.platypus.tableofcontents import TableOfContents
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.pagesizes import A4, A5, A6
from reportlab.pdfgen.canvas import Canvas
from docx import Document
from odf.opendocument import load
from odf.text import P
from ebooklib import epub
from PIL import Image, ImageTk
import os
import re

# Création du dossier output
if not os.path.exists("output"):
    os.mkdir("output")

# Formats disponibles
formats = {
    "A4 (21x29.7 cm)": A4,
    "A5 (14.8x21 cm)": A5,
    "A6 (10.5x14.8 cm)": A6
}

polices = {
    "Times": "Times-Roman",
    "Helvetica": "Helvetica",
    "Courier": "Courier",
    "Georgia": "Times-Roman",
    "Arial": "Helvetica",
    "Comic Sans MS": "Helvetica",
    "Verdana": "Helvetica"
}

interlignes_disponibles = {
    "Simple": 14,
    "1.5": 18,
    "Double": 24
}

# Couleur personnalisée
vert_callia = "#0EDB78"

# Fenêtre principale
app = Window(themename="litera")
app.title("Callia")
app.geometry("1000x800")
app.iconphoto(False, tk.PhotoImage(file="logo.png"))

# Variables de configuration
format_var = tk.StringVar(value="A5 (14.8x21 cm)")
police_var = tk.StringVar(value="Times")
alinea_var = tk.BooleanVar(value=True)
kdp_mode_var = tk.BooleanVar(value=False)
toc_var = tk.BooleanVar(value=True)
stats_var = tk.BooleanVar(value=False)
interligne_var = tk.StringVar(value="1.5")
taille_num_var = tk.StringVar(value="10")
taille_police_var = tk.StringVar(value="12")

# Fonction de génération de PDF
def generate_pdf():
    content = text_field.get("1.0", "end-1c")
    if not content.strip():
        messagebox.showinfo("Erreur", "Le champ est vide.")
        return

    selected_format = formats[format_var.get()]
    selected_police = polices[police_var.get()]
    selected_interligne = interlignes_disponibles[interligne_var.get()]
    taille_police = int(taille_police_var.get())
    taille_num = int(taille_num_var.get())
    width, height = selected_format

    if kdp_mode_var.get():
        marge_gauche, marge_droite = 67, 53
        marge_haut, marge_bas = 67, 67
    else:
        marge_gauche = marge_droite = marge_haut = marge_bas = 72

    pdf_path = filedialog.asksaveasfilename(defaultextension=".pdf", filetypes=[("PDF Files", "*.pdf")], initialfile="mon_livre.pdf")
    if not pdf_path:
        return

    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=selected_format,
        rightMargin=marge_droite,
        leftMargin=marge_gauche,
        topMargin=marge_haut,
        bottomMargin=marge_bas
    )

    styles = getSampleStyleSheet()
    style = ParagraphStyle(
        name="Roman",
        parent=styles["Normal"],
        fontName=selected_police,
        fontSize=taille_police,
        leading=selected_interligne,
        alignment=TA_JUSTIFY,
        spaceAfter=6
    )

    story = []
    if toc_var.get():
        toc = TableOfContents()
        toc.levelStyles = [style]
        story.append(toc)
        story.append(PageBreak())

    for line in content.split("\n"):
        if line.strip():
            txt = "    " + line if alinea_var.get() else line
            story.append(Paragraph(txt, style))
        else:
            story.append(Spacer(1, 12))

    if stats_var.get():
        story.append(PageBreak())
        mots = len(content.split())
        lignes = len(content.splitlines())
        story.append(Paragraph(f"Statistiques : {mots} mots, {lignes} lignes.", style))

    doc.build(story, onFirstPage=lambda c, d: add_page_number(c, d, taille_num), onLaterPages=lambda c, d: add_page_number(c, d, taille_num))
    messagebox.showinfo("Succès", f"PDF généré avec succès dans : {pdf_path}")

# Génération EPUB
def generate_epub():
    content = text_field.get("1.0", "end-1c")
    if not content.strip():
        messagebox.showinfo("Erreur", "Le champ est vide.")
        return

    epub_path = filedialog.asksaveasfilename(defaultextension=".epub", filetypes=[("EPUB Files", "*.epub")], initialfile="mon_livre.epub")
    if not epub_path:
        return

    book = epub.EpubBook()
    book.set_identifier("id123456")
    book.set_title("Mon Livre")
    book.set_language("fr")

    chapter = epub.EpubHtml(title="Chapitre 1", file_name="chap_01.xhtml", lang="fr")
    chapter.content = f'<html><body><p>{content.replace("\n", "<br>")}</p></body></html>'
    book.add_item(chapter)

    book.toc = (epub.Link("chap_01.xhtml", "Chapitre 1", "chap1"),)
    book.spine = ["nav", chapter]
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    epub.write_epub(epub_path, book)
    messagebox.showinfo("Succès", f"EPUB généré avec succès dans : {epub_path}")

# Ajout numérotation

def add_page_number(canvas: Canvas, doc, taille):
    page_num = canvas.getPageNumber()
    canvas.setFont("Helvetica", taille)
    canvas.drawCentredString(doc.pagesize[0] / 2.0, 20, str(page_num))

# Import de fichier texte

def importer_fichier():
    filepath = filedialog.askopenfilename(filetypes=[("Fichiers texte", "*.txt *.docx *.odt")])
    if not filepath:
        return
    try:
        if filepath.endswith(".txt"):
            with open(filepath, "r", encoding="utf-8") as f:
                content = f.read()
        elif filepath.endswith(".docx"):
            doc = Document(filepath)
            content = "\n".join([p.text for p in doc.paragraphs])
        elif filepath.endswith(".odt"):
            odt = load(filepath)
            content = "\n".join([str(p.firstChild.data) for p in odt.getElementsByType(P) if p.firstChild])
        text_field.delete("1.0", tk.END)
        text_field.insert(tk.END, content)
    except Exception as e:
        messagebox.showerror("Erreur", str(e))

# Mise à jour stats en bas de page
def update_stats(event=None):
    content = text_field.get("1.0", "end-1c")
    mots = len(content.split())
    phrases = len(re.findall(r'[.!?]', content))
    temps = max(1, round(mots / 250))
    stats_label.config(text=f"{mots} mots · {phrases} phrases · {temps} min de lecture")

# Interface graphique
main_frame = ttk.Frame(app)
main_frame.place(relx=0.5, rely=0, anchor="n")

logo_img = Image.open("logo.png").resize((50, 50))
logo = ImageTk.PhotoImage(logo_img)
header_frame = ttk.Frame(main_frame)
header_frame.pack(pady=10)

logo_label = ttk.Label(header_frame, image=logo)
logo_label.pack(side="left", padx=5)
title_label = ttk.Label(header_frame, text="Callia — Créateur de livres", font=("Segoe UI", 16, "bold"))
title_label.pack(side="left")

button_frame = ttk.Frame(main_frame)
button_frame.pack(pady=10)

labels = ["Format :", "Police :", "Interligne :", "Taille texte :", "Taille numéros :"]
variables = [format_var, police_var, interligne_var, taille_police_var, taille_num_var]
options = [list(formats.keys()), list(polices.keys()), list(interlignes_disponibles.keys()), None, None]

for i, label in enumerate(labels):
    ttk.Label(button_frame, text=label).grid(row=0, column=2*i, padx=3)
    if options[i]:
        ttk.OptionMenu(button_frame, variables[i], variables[i].get(), *options[i]).grid(row=0, column=2*i+1, padx=3)
    else:
        ttk.Entry(button_frame, textvariable=variables[i], width=4).grid(row=0, column=2*i+1, padx=3)

option_frame = ttk.Frame(main_frame)
option_frame.pack(pady=5)

for widget in [
    ttk.Checkbutton(option_frame, text="Alinéa automatique", variable=alinea_var),
    ttk.Checkbutton(option_frame, text="Format KDP", variable=kdp_mode_var),
    ttk.Checkbutton(option_frame, text="Inclure une table des matières", variable=toc_var),
    ttk.Checkbutton(option_frame, text="Statistiques du texte", variable=stats_var)
]:
    widget.grid(row=0, column=option_frame.grid_size()[0], padx=5)

text_frame = ttk.Frame(main_frame)
text_frame.pack(pady=10)

text_field = tk.Text(text_frame, height=25, width=110)
text_field.pack()
text_field.bind("<KeyRelease>", update_stats)

bottom_frame = ttk.Frame(main_frame)
bottom_frame.pack(pady=10)

for txt, cmd in [("Importer un fichier", importer_fichier), ("Générer le PDF", generate_pdf), ("Générer l'EPUB", generate_epub)]:
    btn = ttk.Button(bottom_frame, text=txt, command=cmd)
    btn.pack(side="left", padx=10)
    btn.configure(bootstyle=f"outline", style=f"success.TButton")
    btn.configure(style=f"{vert_callia}.TButton")

# Statistiques live
stats_label = ttk.Label(app, text="0 mots · 0 phrases · 0 min de lecture", font=("Segoe UI", 10))
stats_label.pack(side="bottom", pady=5)

update_stats()
app.mainloop()
